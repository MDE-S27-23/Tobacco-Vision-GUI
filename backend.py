"""Backend placeholders

Currently this is the ONLY file the GUI calls into

    start_batch()                             -> Tab 1, once before a run
    process_image(image_path)                 -> Tab 1, once per photo
    export_results_to_excel(results, path)    -> Tab 1, "Save Excel file"
    plan_pages(request)                       -> Tab 2, page count for the summary
    render_preview(request, page_index)       -> Tab 2, PNG of one page for the preview
    export_barcodes(request, path)            -> Tab 2, "Save barcodes"

If a function fails, raise an exception with a short, plain message.
"""

import io
import math
import os
import random
import time

# Tab 2: barcodes
#
# The frontend decides WHICH barcodes to print, this file decides HOW they look and where they go
# Every Tab 2 function gets the same request, plain dicts and lists to stay JSON-friendly:
#     {"barcodes": [{"text": "L1042-2026-10-0001"}, ...]}
#
# TODO(backend): replace the 3 placeholders below with the real barcodes and layout
#   - plan_pages: count pages only, keep it instant for typing (called on every keystroke)
#   - render_preview: draw ONE page and return PNG bytes, around 100 dpi is plenty
#   - export_barcodes: write every page to output_path as a PDF
#   - preview and export run in a background thread, keep tkinter out of them
#   - treat the request as read-only, copy it first if you need to change it
#   I came up with from our talk: reportlab draws the PDF, pypdfium2 renders one page of it as the preview PNG
#   One drawing path keeps the preview identical to the print

# File type offered by the "Save barcodes" dialog
BARCODE_FILE_EXTENSION = ".pdf"
BARCODE_FILE_DESCRIPTION = "PDF document"

# Placeholder layout, US Letter upright with barcodes in pairs (two side by side, cut out as one strip)
PAGE_WIDTH_IN = 8.5
PAGE_HEIGHT_IN = 11.0
PAGE_MARGIN_IN = 0.5
PAIRS_PER_ROW = 1
ROWS_PER_PAGE = 15
PLACEHOLDER_DPI = 100


def plan_pages(request):
    """How many pages the barcodes fill. Returns {"page_count": int}"""
    # TODO(backend): real page count from the real layout
    pairs = math.ceil(len(request["barcodes"]) / 2)
    return {"page_count": math.ceil(pairs / (PAIRS_PER_ROW * ROWS_PER_PAGE))}


def render_preview(request, page_index):
    """PNG bytes of one page, page_index starts at 0"""
    # TODO(backend): draw the real page, ideally by rendering that page of the real PDF
    page = _placeholder_page(_split_into_pages(_texts(request))[page_index])
    out = io.BytesIO()
    page.save(out, "PNG")
    return out.getvalue()


def export_barcodes(request, output_path):
    """Write every page to output_path as a PDF"""
    # TODO(backend): write the real PDF with vector barcodes
    pages = [_placeholder_page(pairs) for pairs in _split_into_pages(_texts(request))]
    pages[0].save(
        output_path, "PDF", resolution=PLACEHOLDER_DPI, save_all=True, append_images=pages[1:]
    )


def _texts(request):
    return [barcode["text"] for barcode in request["barcodes"]]


def _split_into_pages(barcode_texts):
    """List of pages, each a list of (left, right) pairs in printing order. right is None for an odd one out"""
    padded = list(barcode_texts) + [None]
    pairs = [(padded[i], padded[i + 1]) for i in range(0, len(barcode_texts), 2)]
    per_page = PAIRS_PER_ROW * ROWS_PER_PAGE
    return [pairs[i : i + per_page] for i in range(0, len(pairs), per_page)]


def _pair_slots():
    """(x, y, width, height) in inches from the top-left paper corner, one per pair position"""
    width = (PAGE_WIDTH_IN - 2 * PAGE_MARGIN_IN) / PAIRS_PER_ROW
    height = (PAGE_HEIGHT_IN - 2 * PAGE_MARGIN_IN) / ROWS_PER_PAGE
    return [
        (PAGE_MARGIN_IN + column * width, PAGE_MARGIN_IN + row * height, width, height)
        for row in range(ROWS_PER_PAGE)
        for column in range(PAIRS_PER_ROW)
    ]


def _placeholder_page(pairs):
    """One page with fake bars and the real text of each barcode"""
    from PIL import Image, ImageDraw  # pip install pillow

    dpi = PLACEHOLDER_DPI
    page = Image.new("L", (round(PAGE_WIDTH_IN * dpi), round(PAGE_HEIGHT_IN * dpi)), 255)
    draw = ImageDraw.Draw(page)
    draw.text((page.width / 2, 25), "PLACEHOLDER - no real barcodes yet", fill=0, anchor="mm")
    for (x, y, width, height), pair in zip(_pair_slots(), pairs):
        left, top = x * dpi, y * dpi
        right, bottom = (x + width) * dpi, (y + height) * dpi
        draw.rectangle([left, top, right, bottom], outline=160)  # cut line
        half = (right - left) / 2
        for index, text in enumerate(pair):
            if text is None:
                continue
            cell_left = left + index * half
            # Random bars seeded by the text to keep each barcode looking the same between renders
            rng = random.Random(text)
            bar_x, bar_right = cell_left + 30, cell_left + half - 30
            while bar_x < bar_right:
                bar = rng.choice((1, 1, 2, 3))
                draw.rectangle([bar_x, top + 8, bar_x + bar - 1, bottom - 22], fill=0)
                bar_x += bar + rng.choice((1, 2, 3))
            draw.text((cell_left + half / 2, bottom - 12), text, fill=0, anchor="mm")
    return page


# Tab 1: tray photos

class SkipPhoto(Exception):
    """Raise from process_image to skip a photo without calling it a failure"""


def start_batch():
    """forget anything remembered from the previous run"""


def process_image(image_path):
    """Analyse ONE tray photo and return its result

    Currently it waits half a second and returns a dummy result
    """
    if not os.path.isfile(image_path):
        raise FileNotFoundError("The photo could not be found.")
    time.sleep(0.5)
    return {"note": "placeholder result"}


def export_results_to_excel(results, output_path):
    """Write the results of a finished batch to an Excel file

    Currently it writes one row per photo with the file name and status
    """
    from openpyxl import Workbook  # pip install openpyxl

    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Results"
    sheet.append(["Filename", "Status", "Result (placeholder)"])
    for result in results:
        if result["ok"]:
            status, detail = "OK", str(result["data"])
        elif result["skipped"]:
            status, detail = "Skipped", result["error"]
        else:
            status, detail = "Failed", result["error"]
        sheet.append([os.path.basename(result["image_path"]), status, detail])
    workbook.save(output_path)
